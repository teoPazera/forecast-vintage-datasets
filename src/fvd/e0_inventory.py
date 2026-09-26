"""Stage E0: inventory of sources and documents, and the vintage calendar.

Run:  python -m fvd.e0_inventory

Writes:
  inventory/sources.csv           provenance of every raw file (appended by fvd.http)
  inventory/documents.csv         every inventoried document, with publication dates
  inventory/vintage_calendar.csv  every forecast vintage with its exact publication date
  inventory/e0_summary.json       counts used in the stage note
"""

from __future__ import annotations

import calendar
import hashlib
import json
import re
import zipfile
from datetime import date

import openpyxl
import pandas as pd
from bs4 import BeautifulSoup

from . import cbo_web, obr_web
from .hofd import series_sheets, vintage_rows
from .http import FetchError, _wayback_lookup, fetch, fetch_page, mark_challenged
from .paths import INVENTORY, RAW_CBO, RAW_OBR, RAW_WEB
from .schemas import SCHEMAS, check_columns, write_schema

CBO_REPO = "US-CBO/eval-projections"
CBO_SHA = "682559ca5800ee7be57a398085fd023a050c7cb9"   # HEAD on 2026-09-26

OBR_DATABASES = {  # label on obr.uk/data/ -> short name
    "Historical official forecasts database": "hofd",
    "Forecast revisions database": "frd",
    "Policy measures database": "pmd",
}

DOC_COLUMNS = ["doc_id", "source", "collection", "doc_type", "title", "url", "landing_url",
               "publication_date", "date_certainty", "date_basis", "related_date",
               "vintage_label", "primary", "notes"]


def doc_id(source: str, url: str) -> str:
    return f"{source.lower()}_{hashlib.sha1(url.encode()).hexdigest()[:10]}"


def month_end(d: date) -> date:
    return date(d.year, d.month, calendar.monthrange(d.year, d.month)[1])


# --- OBR -----------------------------------------------------------------------

def obr_databases() -> dict[str, str]:
    """Download the three OBR databases from the links on obr.uk/data/."""
    s = BeautifulSoup(fetch_page(f"{obr_web.OBR}/data/", RAW_WEB), "lxml")
    found: dict[str, str] = {}
    for a in s.find_all("a", href=True):
        text = a.get_text(" ", strip=True)
        for label, short in OBR_DATABASES.items():
            if text.lower().startswith(label.lower()) and short not in found:
                path = fetch(obr_web._abs(a["href"]), RAW_OBR,
                             note=f"linked from obr.uk/data/ as '{text}'")
                found[short] = path.relative_to(RAW_OBR.parent.parent).as_posix()
    missing = set(OBR_DATABASES.values()) - set(found)
    if missing:
        raise RuntimeError(f"databases not linked from obr.uk/data/: {missing}")
    return found


_TABLE_HINTS = ("charts-and-tables", "charts and tables", "detailed-forecast-tables",
                "detailed forecast tables", "supplementary", "database", "monthly-profiles",
                "monthly profiles", "ready-reckoner", "ready reckoner", "determinants",
                "fiscal-tables", "economy-tables", "-tables", " tables")


_NOT_REPORT = ("summary", "press", "speaking", "slides", "presentation", "restated", "log of",
               "log-of", "correction", "devolved", "briefing", "letter", "presser", "data sources")


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[‐-―-]+", "-", s or "")).strip().lower()


def _classify_efo_file(label: str, url: str, landing: obr_web.Landing) -> tuple[str, bool]:
    """(doc_type, is_primary_report) for a file linked from an EFO landing page.

    The main report is the file whose label is the landing page title, or whose
    label or download slug reads "Economic and fiscal outlook – <Month YYYY>".
    """
    slug = url.rstrip("/").rsplit("/", 1)[-1]
    text = f"{label} {slug}".lower()
    if any(h in text for h in _TABLE_HINTS):
        return "forecast_tables", False
    if any(h in text for h in _NOT_REPORT):
        return "other", False
    my = _month_year(landing.title).lower()
    lab = _norm(label)
    if my and (lab == _norm(landing.title)
               or lab in (f"economic and fiscal outlook - {my}", f"{my} economic and fiscal outlook",
                          f"budget forecast - {my}", f"pre-budget forecast - {my}")
               or re.fullmatch(rf"economic-(and-)?fiscal-outlook-{my.replace(' ', '-')}(-\d+)?", slug)):
        return "forecast_narrative", True
    return "other", False


def obr_documents() -> tuple[list[dict], list[obr_web.Landing], list[obr_web.Landing]]:
    rows: list[dict] = []
    efo = obr_web.walk_chain(f"{obr_web.OBR}/efo/economic-and-fiscal-outlook-march-2026/",
                             "Previous forecast")
    fer = obr_web.walk_chain(f"{obr_web.OBR}/fer/forecast-evaluation-report-june-2026/",
                             "Previous report")
    for page in efo:
        vlabel = _month_year(page.title)
        for label, url in page.files:
            dtype, primary = _classify_efo_file(label, url, page)
            rows.append(_obr_row("efo", dtype, label or url.rsplit("/", 2)[-2], url, page,
                                 vlabel, primary))
    for page in fer:
        for label, url in page.files:
            slug = url.rstrip("/").rsplit("/", 1)[-1]
            primary = slug == page.url.rstrip("/").rsplit("/", 1)[-1]
            rows.append(_obr_row("fer", "post_hoc_evaluation", label or slug, url, page, "", primary))
    rows += obr_commentary()
    return rows, efo, fer


def _month_year(title: str) -> str:
    m = re.search(rf"({obr_web_months()})\s+(\d{{4}})", title)
    if m:
        return f"{m.group(1)} {m.group(2)}"
    m = re.search(r"(\d{4})", title)   # "June Budget 2010" -> June 2010
    mm = re.search(rf"({obr_web_months()})", title)
    return f"{mm.group(1)} {m.group(1)}" if m and mm else ""


def obr_web_months() -> str:
    return "|".join(calendar.month_name[1:])


def _obr_row(collection, dtype, title, url, page, vlabel, primary) -> dict:
    return {
        "doc_id": doc_id("OBR", url), "source": "OBR", "collection": collection,
        "doc_type": dtype, "title": title, "url": url, "landing_url": page.url,
        "publication_date": page.publication_date.isoformat() if page.publication_date else "",
        "date_certainty": "exact" if page.publication_date else "unknown",
        "date_basis": "landing page <time class=date>" if page.publication_date else "",
        "related_date": "", "vintage_label": vlabel, "primary": primary, "notes": "",
    }


def obr_commentary() -> list[dict]:
    """Monthly commentary on the public sector finances, from the yearly archive pages."""
    from .http import wayback_captures

    base = f"{obr_web.OBR}/monthly-public-finances-briefing/"
    years = [u for _, u in obr_web.listing_links(base, r"/monthly-public-finances-briefing/.+")]
    # The archived index page predates the 2023-24 to 2025-26 year pages; those
    # pages exist in the archive, so add every year page it holds.
    for c in wayback_captures(base.split("://")[1], match="prefix", collapse="urlkey"):
        u = obr_web._abs(c["original"].split("?")[0])
        if re.search(r"/monthly-public-finances-briefing/[^/]*\d{4}-\d{2}/?$", u) and \
                u.rstrip("/") not in {y.rstrip("/") for y in years}:
            years.append(u if u.endswith("/") else u + "/")
    rows = []
    seen = set()
    for yurl in years:
        for label, url in obr_web.listing_links(yurl, r"/download/|/docs/"):
            if "commentary" not in label.lower() or url in seen:
                continue
            seen.add(url)
            m = re.search(rf"({obr_web_months()})\s+(\d{{4}})", label)
            if not m:
                # 2011-12 and 2012-13 pages label every link "Commentary on the public
                # sector finances release"; the month is read from the file in E7.
                rows.append({
                    "doc_id": doc_id("OBR", url), "source": "OBR", "collection": "psf_commentary",
                    "doc_type": "in_period_commentary", "title": label, "url": url,
                    "landing_url": yurl, "publication_date": "", "date_certainty": "unknown",
                    "date_basis": "listing label has no month; resolve from the file in E7",
                    "related_date": "", "vintage_label": "", "primary": False,
                    "notes": f"year page {yurl.rstrip('/').rsplit('/', 1)[-1]}",
                })
                continue
            data_month = date(int(m.group(2)), list(calendar.month_name).index(m.group(1)), 1)
            # The commentary follows the ONS release for that data month, which comes
            # out during the next month; only the month is known from the listing.
            pub_month = date(data_month.year + (data_month.month == 12),
                             data_month.month % 12 + 1, 1)
            rows.append({
                "doc_id": doc_id("OBR", url), "source": "OBR", "collection": "psf_commentary",
                "doc_type": "in_period_commentary", "title": label, "url": url,
                "landing_url": yurl, "publication_date": month_end(pub_month).isoformat(),
                "date_certainty": "month_only",
                "date_basis": "data month + 1 (ONS release month); last day of month (D5)",
                "related_date": "", "vintage_label": "", "primary": False,
                "notes": f"data month {data_month:%Y-%m}",
            })
    return rows


# --- CBO -----------------------------------------------------------------------

def cbo_repository() -> str:
    path = fetch(f"https://api.github.com/repos/{CBO_REPO}/zipball/{CBO_SHA}", RAW_CBO,
                 name=f"eval-projections-{CBO_SHA[:7]}.zip", note=f"{CBO_REPO} commit {CBO_SHA}")
    out = RAW_CBO / f"eval-projections-{CBO_SHA[:7]}"
    if not out.exists():
        with zipfile.ZipFile(path) as z:
            root = z.namelist()[0].split("/")[0]
            z.extractall(RAW_CBO / "_extract")
        (RAW_CBO / "_extract" / root).rename(out)
        (RAW_CBO / "_extract").rmdir()
    return out.relative_to(RAW_CBO.parent.parent).as_posix()


_CBO_SECTIONS = {  # recurring-reports section -> (collection, doc_type)
    "Budget and Economic Outlook and Updates": ("cbo_beo", "forecast_narrative"),
    "Analysis of the President's Budget": ("cbo_apb", "forecast_narrative"),
    "Accuracy of CBO’s Baseline Projections": ("cbo_accuracy", "post_hoc_evaluation"),
    "Accuracy of CBO's Baseline Projections": ("cbo_accuracy", "post_hoc_evaluation"),
    "Monthly Budget Review": ("cbo_mbr", "in_period_commentary"),
}

# Evaluation reports named in the eval-projections README.
CBO_EVALUATIONS = [
    ("An Evaluation of CBO's Projections of Deficits and Debt From 1984 to 2023", "https://www.cbo.gov/publication/60664"),
    ("An Evaluation of CBO's Projections of Outlays from 1984 to 2021", "https://www.cbo.gov/publication/58613"),
    ("An Evaluation of CBO's Past Revenue Projections", "https://www.cbo.gov/publication/56499"),
]


def _cbo_row(item: cbo_web.Listed, collection: str, dtype: str, basis: str, notes="") -> dict:
    if item.pub_date:
        pub, cert = item.pub_date.isoformat(), "exact"
    else:
        pub, cert = month_end(date.fromisoformat(item.month + "-01")).isoformat(), "month_only"
    return {
        "doc_id": doc_id("CBO", item.url), "source": "CBO", "collection": collection,
        "doc_type": dtype, "title": item.title, "url": item.url, "landing_url": item.url,
        "publication_date": pub, "date_certainty": cert,
        "date_basis": basis + ("" if cert == "exact" else "; last day of month (D5)"),
        "related_date": "", "vintage_label": "", "primary": dtype == "forecast_narrative",
        "notes": notes,
    }


def cbo_documents() -> list[dict]:
    rows = []
    for item in cbo_web.recurring_reports():
        if item.section not in _CBO_SECTIONS:
            continue
        coll, dtype = _CBO_SECTIONS[item.section]
        basis = "cbo.gov recurring-reports listing <time datetime>"
        if item.day_is_placeholder:
            basis += " (day 01 treated as month-level)"
        rows.append(_cbo_row(item, coll, dtype, basis))

    # Pre-2000 reports come from the pre-2012 site's subject listings. Testimony on
    # an outlook is kept as dated CBO text (doc_type "other").
    old = {cat: cbo_web.old_by_subject(cat) for cat in cbo_web.OLD_CATEGORIES}
    for cat, items in old.items():
        coll = {0: "cbo_beo", 1: "cbo_apb", 35: "cbo_mbr"}[cat]
        for item in items:
            testimony = item.title.lower().startswith("testimony")
            # From 2000 the current site lists the reports; testimony only comes from here.
            if int(item.month[:4]) >= 2000 and not testimony:
                continue
            if cat == 0 and not testimony and not re.search(r"outlook", item.title, re.I):
                continue   # the subject also lists unrelated studies
            dtype = ("other" if testimony else
                     {0: "forecast_narrative", 1: "forecast_narrative", 35: "in_period_commentary"}[cat])
            rows.append(_cbo_row(item, "cbo_testimony" if testimony else coll, dtype,
                                 "pre-2012 cbo.gov subject listing (Wayback capture)"))
    for title, url in CBO_EVALUATIONS:
        rows.append({"doc_id": doc_id("CBO", url), "source": "CBO", "collection": "cbo_evaluation",
                     "doc_type": "post_hoc_evaluation", "title": title, "url": url,
                     "landing_url": url, "publication_date": "", "date_certainty": "unknown",
                     "date_basis": "named in eval-projections README; date resolved in E7",
                     "related_date": "", "vintage_label": "", "primary": False, "notes": ""})
    return rows


# --- vintage calendar ----------------------------------------------------------

def hofd_labels(hofd_path) -> pd.DataFrame:
    """Every vintage label in the unsuffixed HOFD sheets, in sheet order."""
    wb = openpyxl.load_workbook(hofd_path, read_only=True, data_only=True)
    recs = []
    for name in series_sheets(wb.sheetnames):
        rows = list(wb[name].iter_rows(values_only=True))
        for i, lab in vintage_rows(rows):
            recs.append({"sheet": name, "row": i + 1, "raw": lab.raw, "label": lab.label,
                         "kind": lab.kind, "memo_type": lab.memo_type,
                         "vintage_label": lab.vintage_label, "footnote": lab.footnote})
    return pd.DataFrame(recs)


def _label_date(label: str) -> date:
    mon, yr = label.split()
    return date(int(yr), list(calendar.month_name).index(mon), 1)


def obr_calendar(labels: pd.DataFrame, efo: list[obr_web.Landing], docs: list[dict]) -> list[dict]:
    by_label: dict[str, list[obr_web.Landing]] = {}
    for p in efo:
        by_label.setdefault(_month_year(p.title), []).append(p)
    primary = {d["landing_url"]: d["doc_id"] for d in docs
               if d["collection"] == "efo" and d["primary"]}
    counts = labels.groupby("vintage_label").sheet.nunique()
    first_seen = labels.drop_duplicates("vintage_label")
    out = []
    for _, r in first_seen.iterrows():
        d0 = _label_date(r.label)
        forecaster = "OBR" if d0 >= date(2010, 6, 1) else "HM Treasury"
        rec = {"source": "OBR", "vintage_id": "", "label": r.vintage_label,
               "forecaster": forecaster, "publication_date": "", "date_certainty": "",
               "date_basis": "", "landing_url": "", "primary_document_id": "",
               "flags": "", "n_sheets": int(counts[r.vintage_label]), "notes": ""}
        flags = []
        if r.kind == "memo":
            flags.append(f"memo:{r.memo_type}")
            rec["date_certainty"] = "unknown"
            rec["notes"] = "memo vintage; publication date to be established in E1"
        elif forecaster == "OBR":
            pages = by_label.get(r.label, [])
            if r.label == "June 2010":
                # Two June 2010 publications: the Pre-Budget forecast (14 June) and the
                # June Budget EFO (22 June). The HOFD has one June 2010 row.
                pages = [p for p in pages if "budget-2010" in p.url and "pre-budget" not in p.url]
                flags.append("june2010_budget_assumed")
                rec["notes"] = ("HOFD 'June 2010' mapped to the June Budget 2010 EFO (22 June), "
                                "not the Pre-Budget forecast (14 June); check values in E2")
            if len(pages) == 1 and pages[0].publication_date:
                p = pages[0]
                rec.update(publication_date=p.publication_date.isoformat(),
                           date_certainty="exact", date_basis="EFO landing page",
                           landing_url=p.url, primary_document_id=primary.get(p.url, ""))
            else:
                rec.update(date_certainty="unknown",
                           notes=f"{len(pages)} EFO landing pages match this label")
        else:
            rec.update(publication_date=month_end(d0).isoformat(), date_certainty="month_only",
                       date_basis="HM Treasury vintage label; last day of month (D5)")
            flags.append("hmt")
        rec["flags"] = ";".join(flags)
        stamp = d0.strftime("%Y-%m")
        rec["vintage_id"] = f"obr_{stamp}" + (f"_{r.memo_type}" if r.kind == "memo" else "")
        if forecaster == "HM Treasury":
            rec["vintage_id"] = f"hmt_{stamp}"
        out.append(rec)
    return out


def cbo_calendar(repo_dir: str, docs: list[dict]) -> list[dict]:
    b = pd.read_csv(f"{repo_dir}/input_data/baselines.csv")
    flags = b.groupby("baseline_date").agg(spring=("Spring_flag", "max"),
                                           winter=("Winter_flag", "max"),
                                           components=("component", lambda s: ",".join(sorted(set(s)))))
    cand = [d for d in docs if d["source"] == "CBO" and d["collection"] in ("cbo_beo", "cbo_apb")]
    testimony = [d for d in docs if d["collection"] == "cbo_testimony"]
    out = []
    for bdate, f in flags.iterrows():
        month = bdate[:7]
        same = [d for d in cand if d["publication_date"][:7] == month]
        exact = sorted([d for d in same if d["date_certainty"] == "exact"],
                       key=lambda d: (d["collection"] != "cbo_beo", d["publication_date"]))
        # On the pre-2012 listing, reports carry a month and statements a full date. If a
        # month-only report exists in the same month, an exact-dated item is a statement
        # about it, not the report: keep the report month-only and the day as related.
        old_month_only = [d for d in same if "doc.cfm" in d["url"] and d["date_certainty"] != "exact"]
        demoted = [d for d in exact if "doc.cfm" in d["url"]] if old_month_only else []
        exact = [d for d in exact if d not in demoted]
        fl = [x for x, on in (("spring", f.spring), ("winter", f.winter)) if on]
        rec = {"source": "CBO", "vintage_id": f"cbo_{month}", "label": month,
               "forecaster": "CBO", "flags": ";".join(fl) or "no_season_flag",
               "components": f.components, "n_candidates": len(same), "related_date": ""}
        if exact:
            d = exact[0]
            rec.update(publication_date=d["publication_date"], date_certainty="exact",
                       date_basis=d["date_basis"], primary_document_id=d["doc_id"],
                       document_url=d["url"])
            if len(exact) > 1:
                rec["notes"] = "several exact-dated reports in the month: " + ", ".join(
                    f"{x['collection']} {x['publication_date']}" for x in exact)
            else:
                rec["notes"] = ""
        else:
            d = old_month_only[0] if old_month_only else (same[0] if same else None)
            rec.update(publication_date=month_end(date.fromisoformat(month + "-01")).isoformat(),
                       date_certainty="month_only", date_basis="baseline month; last day of month (D5)",
                       primary_document_id=d["doc_id"] if d else "",
                       document_url=d["url"] if d else "",
                       notes="" if d else "no report listed in the baseline month")
            tdates = sorted([t["publication_date"] for t in testimony
                             if t["publication_date"][:7] == month and t["date_certainty"] == "exact"]
                            + [x["publication_date"] for x in demoted])
            if tdates:
                rec["related_date"] = tdates[0]
                rec["notes"] = (rec["notes"] + "; " if rec["notes"] else "") + \
                    "earliest dated CBO statement or testimony on the outlook that month: " + tdates[0]
        out.append(rec)
    return out


# --- availability check ----------------------------------------------------------

_REPORT_WORDS = {"efo": ("efo", "outlook", "forecast", "budget"), "fer": ("fer", "evaluation")}
# The two June 2010 publications share every generic phrase; their title pages differ.
_TITLE_PAGE_RULES = {
    "pre-budget-forecast-june-2010": lambda p1: "pre-budget forecast" in p1,
    "budget-2010": lambda p1: "budget forecast" in p1 and "pre-budget" not in p1,
}
_REPORT_TITLES = {"efo": ("economic and fiscal outlook", "budget forecast"),
                  "fer": ("forecast evaluation report",)}
_ALT_EXCLUDE = ("summary", "press", "_pn", "-pn", "pn.", "presser", "speaking", "slides",
                "presentation", "briefing", "log", "chart", "table", "supplementary", "annex",
                "correction", "box", "overview", "letter", "commentary")


def _first_pages_text(path, n=3) -> str:
    import pymupdf as fitz

    try:
        with fitz.open(stream=path.read_bytes(), filetype="pdf") as pdf:
            return " ".join(pdf[i].get_text() for i in range(min(n, pdf.page_count)))
    except Exception:
        return ""


def _page_count(path) -> int:
    import pymupdf as fitz

    try:
        with fitz.open(stream=path.read_bytes(), filetype="pdf") as pdf:
            return pdf.page_count
    except Exception:
        return 0


def _recover_report(d: dict, month_year: str) -> str:
    """Find an archived copy of a report whose own link was never archived.

    Candidates are archived PDFs on obr.uk or the OBR's former domains whose
    file name names the publication's month and year; a candidate is accepted
    only if its first pages carry the report title and the month and year.
    """
    from .http import fetch_capture
    from .paths import RAW_DOCS

    month, year = month_year.split()
    rule = _TITLE_PAGE_RULES.get(d["landing_url"].rstrip("/").rsplit("/", 1)[-1])

    def is_report(path) -> bool:
        # A full report runs to dozens of pages; annexes and notes are shorter.
        if _page_count(path) < 25:
            return False
        # six pages: some covers are images with no text layer
        text = re.sub(r"\s+", " ", _first_pages_text(path, n=6)).lower()
        ok = (any(t in text for t in _REPORT_TITLES[d["collection"]])
              and month.lower() in text and year in text)
        if ok and rule is not None:
            ok = rule(re.sub(r"\s+", " ", _first_pages_text(path, n=1)).lower())
        return ok

    # FER file names often carry only the year, and some early files only the month
    # ("junebudget_annexc.pdf"); the title-page check still needs both.
    words = _REPORT_WORDS[d["collection"]]
    passes = [obr_web.report_candidates(month, int(year), words, _ALT_EXCLUDE,
                                        require_month=d["collection"] != "fer")[:8],
              obr_web.report_candidates(month, int(year), words,
                                        tuple(x for x in _ALT_EXCLUDE if x != "annex"),
                                        require_year=False)[:8]]
    tried = set()
    for candidates in passes:
        for c in candidates:
            if c["original"] in tried:
                continue
            tried.add(c["original"])
            try:
                path = fetch_capture(d["url"], c["original"], RAW_DOCS / "obr", verify=is_report,
                                     note="source link not archived; same report archived "
                                          "elsewhere, title page verified")
            except FetchError:
                continue
            if path is not None:
                return c["original"]
    return ""


def check_available(docs: list[dict], landings: dict[str, obr_web.Landing]) -> None:
    """Mark whether each primary EFO and FER report can be retrieved."""
    mark_challenged("obr.uk")
    for d in docs:
        if d["source"] != "OBR" or not d["primary"] or d["collection"] not in ("efo", "fer"):
            continue
        try:
            ok = _wayback_lookup(d["url"], "latest") is not None
        except FetchError:
            ok = False
        if ok:
            status = "report capture available"
        else:
            alt = _recover_report(d, _month_year(landings[d["landing_url"]].title))
            status = (f"report link not archived; same report archived at {alt} (title page verified)"
                      if alt else "REPORT NOT AVAILABLE")
        d["notes"] = (d["notes"] + "; " if d["notes"] else "") + status


# --- main ------------------------------------------------------------------------

def main() -> None:
    INVENTORY.mkdir(parents=True, exist_ok=True)
    dbs = obr_databases()
    repo = cbo_repository()

    obr_docs, efo, fer = obr_documents()
    check_available(obr_docs, {p.url: p for p in efo + fer})
    cbo_docs = cbo_documents()
    docs = obr_docs + cbo_docs
    ddf = pd.DataFrame(docs, columns=DOC_COLUMNS).drop_duplicates("doc_id")
    ddf.to_csv(INVENTORY / "documents.csv", index=False)

    labels = hofd_labels(RAW_OBR.parent.parent / dbs["hofd"])
    labels.to_csv(INVENTORY / "hofd_vintage_labels.csv", index=False)
    cal = pd.DataFrame(obr_calendar(labels, efo, docs) + cbo_calendar(repo, docs))
    cols = [c[0] for c in SCHEMAS["inventory/vintage_calendar"]["columns"]]
    cal = cal.reindex(columns=cols)
    cal["_d"] = cal.vintage_id.str.extract(r"_(\d{4}-\d{2})")[0]
    cal = cal.sort_values(["source", "_d", "vintage_id"]).drop(columns="_d")
    cal.to_csv(INVENTORY / "vintage_calendar.csv", index=False)

    for table, df in [("inventory/documents", ddf), ("inventory/vintage_calendar", cal),
                      ("inventory/hofd_vintage_labels", labels),
                      ("inventory/sources", pd.read_csv(INVENTORY / "sources.csv"))]:
        check_columns(table, df)
        write_schema(table, INVENTORY.parent)

    obr_v = cal[(cal.source == "OBR") & (cal.forecaster == "OBR") & ~cal["flags"].str.contains("memo")]
    cbo_v = cal[cal.source == "CBO"]
    efo_primary = ddf[(ddf.collection == "efo") & (ddf.primary)]
    summary = {
        "databases": dbs, "cbo_repo": repo,
        "efo_landing_pages": len(efo), "fer_landing_pages": len(fer),
        "documents_by_collection": ddf.groupby("collection").size().to_dict(),
        "documents_by_certainty": ddf.groupby("date_certainty").size().to_dict(),
        "obr_vintages": len(obr_v),
        "obr_vintages_exact": int((obr_v.date_certainty == "exact").sum()),
        "hmt_vintages": int((cal.forecaster == "HM Treasury").sum()),
        "memo_vintages": int(cal["flags"].str.contains("memo").sum()),
        "efo_primary_reports": len(efo_primary),
        "efo_reports_unavailable": int(efo_primary.notes.str.contains("NOT AVAILABLE").sum()),
        "cbo_baselines": len(cbo_v),
        "cbo_baselines_exact": int((cbo_v.date_certainty == "exact").sum()),
        "cbo_exact_share": round(float((cbo_v.date_certainty == "exact").mean()), 3),
    }
    (INVENTORY / "e0_summary.json").write_text(json.dumps(summary, indent=2, default=str))
    print(json.dumps(summary, indent=2, default=str))


if __name__ == "__main__":
    main()
